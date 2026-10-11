#!/usr/bin/env python3
"""Load the MASTER DATABASE workbook into the Postgres ``companies`` table.

Reads the "Master Report" sheet (Account Name + technology columns) and upserts
each row keyed by its normalized company name. Repeated account variants that
share a normalized base name are merged, and per-category technology lists are
combined.

Usage:
    python ingest_master.py "/path/to/MASTER DATABASE Oct 2026.xlsx" [--sheet "Master Report"] [--create]
"""
import argparse
import sys

from db import (
    MASTER_SOFTWARE_COLUMNS, _is_noise, extract_domain, init_schema,
    normalize_company_name, split_multi, upsert_companies,
)

KNOWN_COLUMNS = {
    'Account ID': 'account_id',
    'Account Owner': 'account_owner',
    'Account Name': 'account_name',
    'Billing State/Province': 'billing_state',
    'Type': 'account_type',
    'EPM': 'EPM',
    'ERP': 'ERP',
    'CRM': 'CRM',
    'Data Tools': 'Data Tools',
    'Business Intelligence Software': 'Business Intelligence Software',
    'Cloud': 'Cloud',
    'Other': 'Other',
    'Contacts': 'contacts',
    'Contact Titles': 'contact_titles',
    'Source Links': 'source_links',
    'LinkedIn URL': 'linkedin_url',
    'Website URL': 'website_url',
    'Industry': 'industry',
    'Location': 'location',
    'Partner': 'partner',
    'Notes on Implementation': 'notes',
    'Data Sources': 'data_sources',
    'Source Account Name(s)': 'source_account_names',
    'Match Method': 'match_method',
}


def col_index(header):
    out = {}
    for i, h in enumerate(header):
        key = (h or '').strip()
        if key in KNOWN_COLUMNS and KNOWN_COLUMNS[key] not in out:
            out[KNOWN_COLUMNS[key]] = i
    return out


def build_companies(rows, idx):
    merged = {}

    def get(r, k):
        i = idx.get(k)
        v = str(r[i]).strip() if i is not None and i < len(r) else ''
        return '' if _is_noise(v) else v

    for r in rows:
        account_name = get(r, 'account_name')
        if not account_name:
            continue
        norm = normalize_company_name(account_name)
        rec = merged.setdefault(
            norm,
            {
                'normalized_name': norm,
                'account_name': account_name,
                'software': {},
                'raw': {},
            },
        )
        if len(account_name) >= len(rec['account_name']):
            rec['account_name'] = account_name
        for field, mapped in (('account_id', 'account_id'),
                              ('linkedin_url', 'linkedin_url'),
                              ('website_url', 'website_url'),
                              ('industry', 'industry'),
                              ('location', 'location'),
                              ('partner', 'partner'),
                              ('account_type', 'account_type')):
            v = get(r, mapped)
            if v and not rec.get(mapped):
                rec[mapped] = v
        for col, category in MASTER_SOFTWARE_COLUMNS.items():
            vals = split_multi(get(r, col)) if col != 'Other' else _other_values(get(r, col))
            if vals:
                rec['software'].setdefault(category, [])
                for v in vals:
                    if v not in rec['software'][category]:
                        rec['software'][category].append(v)
        # raw / verbatim row
        raw = {}
        for k, i in idx.items():
            v = str(r[i]).strip() if i < len(r) else ''
            if v and not _is_noise(v):
                raw[k] = v
        if raw:
            rec['raw'].setdefault(account_name, raw)

    companies = list(merged.values())
    for rec in companies:
        rec['domain'] = rec.get('domain') or extract_domain(rec.get('website_url'))
        rec['enrichment_source'] = 'db'
    for c in companies:
        c['raw'] = c['raw'] or None
    return companies


def _other_values(v):
    """Filter obvious non-technology filler from the Other column."""
    noise = {'not listed', 'other tools are not publicly listed',
             'all technology not publically disclosed', 'na', 'n/a', 'unknown'}
    return [x for x in split_multi(v) if x.lower() not in noise]


def main():
    ap = argparse.ArgumentParser(description='Load MASTER DATABASE into Postgres companies table.')
    ap.add_argument('workbook', help='Path to the .xlsx master database')
    ap.add_argument('--sheet', default='Master Report', help='Sheet name (default: Master Report)')
    ap.add_argument('--create', action='store_true', help='Create schema before loading')
    ap.add_argument('--max-rows', type=int, default=None, help='Only load the first N data rows')
    args = ap.parse_args()

    from openpyxl import load_workbook

    wb = load_workbook(args.workbook, read_only=True, data_only=True)
    if args.sheet not in wb.sheetnames:
        print(f'Sheet "{args.sheet}" not found. Available: {wb.sheetnames}')
        sys.exit(1)
    ws = wb[args.sheet]
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() for h in (next(it, None) or [])]
    idx = col_index(header)
    if 'account_name' not in idx:
        print('Could not find "Account Name" column')
        sys.exit(1)

    rows = []
    for r in it:
        rows.append(r)
        if args.max_rows and len(rows) >= args.max_rows:
            break
    wb.close()

    if args.create:
        init_schema()

    companies = build_companies(rows, idx)
    print(f'Parsed {len(rows)} master rows -> {len(companies)} distinct normalized companies')
    written = upsert_companies(companies)
    print(f'Upserted {written} companies into Postgres')


if __name__ == '__main__':
    main()