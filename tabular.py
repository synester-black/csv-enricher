"""Read CSV / XLSX / XLS into a list of rows regardless of the source format."""
import csv
import os


def _clean(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_csv(path):
    with open(path, 'r', encoding='utf-8-sig', newline='') as f:
        reader = csv.reader(f)
        header = next(reader, None) or []
        rows = [row for row in reader]
    return [_clean(h) for h in header], [[_clean(x) for x in r] for r in rows]


def read_xlsx(path):
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    header = next(it, None)
    if header is None:
        wb.close()
        return [], []
    header = [_clean(h) for h in header]
    rows = [[_clean(x) for x in r] for r in it]
    wb.close()
    return header, rows


def read_xls(path):
    import xlrd

    book = xlrd.open_workbook(path)
    sheet = book.sheet_by_index(0)
    header = [_clean(c.value) for c in sheet.row(0)] if sheet.nrows else []
    rows = [
        [_clean(sheet.cell_value(r, c)) for c in range(sheet.ncols)]
        for r in range(1, sheet.nrows)
    ]
    return header, rows


_READERS = {
    '.csv': read_csv,
    '.xlsx': read_xlsx,
    '.xlsm': read_xlsx,
    '.xls': read_xls,
}

SUPPORTED_EXTENSIONS = tuple(_READERS.keys())


def read_tabular(path):
    ext = os.path.splitext(str(path))[1].lower()
    if ext not in _READERS:
        raise ValueError(f'Unsupported file type: {ext}')
    return _READERS[ext](path)