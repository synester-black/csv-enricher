"""Export processed results in the "Vicit Data Request" workbook format.

66 columns (A..BN), a single header row, one row per contact/lead. Account-level
fields come from the master database; contact fields come from the uploaded row;
technology, intent and tier are enriched. Empty cells are left blank and
Contact Source / Contact Campaign Type default to "uncontacted".
"""
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Single header row (A..BN).
COLUMN_NAMES = [
    "Account ID", "Contact ID", "Account Owner ID", "Contact Owner ID",
    "Account Owner Name", "Contact Owner Name", "Account Client Owner",
    "Contact Client Owner", "Account Name", "Account Phone", "Employee Count",
    "Website", "Company LinkedIn URL", "Industry", "Other Industry",
    "Parent Account", "Account Type", "X-Factor / Target Flag",
    "Account Description", "Company Overview / Pitch Intel", "Product Focus",
    "Intent Data", "EPM Tool", "ERP Tool", "Data Tools",
    "Business Intelligence Software", "CRM Platform", "Cloud", "Other",
    "Billing Street", "Billing City", "Billing County/State", "Billing Postcode",
    "Billing Country", "Shipping Street", "Shipping City", "Shipping County/State",
    "Shipping Postcode", "Shipping Country", "Salutation", "First Name",
    "Last Name", "Title", "Phone", "Mobile", "Other Phone", "Email",
    "Personal Email", "LinkedIn Profile URL", "Function", "Tier", "Lead Score",
    "Contact Description", "Contact Intel", "Contact Campaign Type",
    "Campaign Name", "Contact Source", "Previous Contact Owner",
    "Previous Lead Score", "Last Activity", "Source Links", "Partner",
    "Technology Information", "Notes on Implementation", "Data Sources",
    "Match Method",
]

DEFAULT_LEAD_SOURCE = "uncontacted"

# Insurance of column indices (0-based) for clarity.
IDX = {name: i for i, name in enumerate(COLUMN_NAMES)}

# db software column -> target column index
SOFTWARE_COLUMNS = {
    "epm_tool": IDX["EPM Tool"],
    "erp_tool": IDX["ERP Tool"],
    "data_tools": IDX["Data Tools"],
    "business_intelligence_software": IDX["Business Intelligence Software"],
    "crm_platform": IDX["CRM Platform"],
    "cloud": IDX["Cloud"],
    "other": IDX["Other"],
}

_FINANCE_HINTS = ('finance', 'fp&a', 'fpa', 'financial', 'cfo', 'accounting',
                  'treasury', 'controller', 'commercial', 'planning', 'reporting',
                  'audit', 'epm')


def _join_multi(v):
    if not v:
        return ""
    if isinstance(v, (list, tuple)):
        return "; ".join(str(x).strip() for x in v if str(x).strip())
    return str(v).strip()


def _function_for(title):
    t = (title or '').lower()
    if any(k in t for k in _FINANCE_HINTS):
        return 'Finance'
    return ''


def build_template_row(account, contact=None, enrichment=None):
    """Return a 66-cell row.

    ``account``   - dict of company columns (account_name, website_url, epm_tool, ...)
    ``contact``   - dict of contact fields (first_name, last_name, title, email, ...)
    ``enrichment``- dict with optional 'software' (column->value), 'intent',
                    'tier', 'source'.
    """
    account = account or {}
    contact = contact or {}
    enrichment = enrichment or {}
    raw = account.get('raw') or {}
    sw = {**{c: account.get(c) for c in SOFTWARE_COLUMNS},
          **(enrichment.get('software') or {})}
    intent = enrichment.get('intent') or account.get('intent_data') or ''
    row = [""] * len(COLUMN_NAMES)

    def setc(name, val):
        if val not in (None, '', [], {}):
            row[IDX[name]] = val

    # --- Account block ---
    setc("Account ID", raw.get('Account ID') or account.get('account_id'))
    setc("Account Owner Name", raw.get('Account Owner'))
    setc("Account Name", account.get('account_name'))
    setc("Employee Count", account.get('employees'))
    setc("Website", account.get('website_url'))
    setc("Company LinkedIn URL", account.get('linkedin_url'))
    setc("Industry", account.get('industry'))
    setc("Account Type", account.get('account_type'))
    setc("Intent Data", intent if isinstance(intent, str) else _join_multi(intent))

    # --- Technology ---
    for col, idx in SOFTWARE_COLUMNS.items():
        setc(COLUMN_NAMES[idx], _join_multi(sw.get(col)))
    tech_dump = {c: _join_multi(sw.get(c)) for c in SOFTWARE_COLUMNS if _join_multi(sw.get(c))}
    setc("Technology Information", json.dumps(tech_dump) if tech_dump else "")

    # --- Billing / location ---
    setc("Billing City", account.get('location') or contact.get('city'))
    setc("Billing Country", contact.get('country'))

    # --- Contact block ---
    setc("Salutation", contact.get('salutation'))
    setc("First Name", contact.get('first_name'))
    setc("Last Name", contact.get('last_name'))
    setc("Title", contact.get('title'))
    setc("Phone", contact.get('phone'))
    setc("Mobile", contact.get('mobile'))
    setc("Other Phone", contact.get('mobile_other'))
    setc("Email", contact.get('email'))
    setc("Personal Email", contact.get('personal_email'))
    setc("LinkedIn Profile URL", contact.get('linkedin_url'))
    setc("Function", contact.get('function') or _function_for(contact.get('title')))
    setc("Tier", enrichment.get('tier') or contact.get('tier'))
    setc("Contact Campaign Type", DEFAULT_LEAD_SOURCE)
    setc("Contact Source", DEFAULT_LEAD_SOURCE)

    # --- Audit / provenance ---
    setc("Last Activity", raw.get('Last Activity'))
    setc("Source Links", raw.get('Source Links'))
    setc("Partner", account.get('partner'))
    setc("Notes on Implementation", raw.get('Notes on Implementation'))
    setc("Data Sources", enrichment.get('source') or "")
    setc("Match Method", raw.get('Match Method') or account.get('match_method') or "")

    return row


def write_template_xlsx(data_rows, out_path, sheet_name="Data Request"):
    """Write a single-header-row workbook (A..BN)."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    for i, name in enumerate(COLUMN_NAMES, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = Font(bold=True)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.fill = PatternFill("solid", fgColor="DCE6F1")

    for r_idx, row in enumerate(data_rows, start=2):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMN_NAMES))}1"
    for i in range(1, len(COLUMN_NAMES) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 18

    wb.save(out_path)
    return out_path